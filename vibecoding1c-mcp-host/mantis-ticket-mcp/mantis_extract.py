"""Bounded, isolated text extraction. Never render or execute document content."""
from __future__ import annotations

import io
import json
import os
import sys
import zipfile


MAX_INPUT = 5 * 1024 * 1024
MAX_TEXT = 100_000
MAX_ZIP_UNCOMPRESSED = 50 * 1024 * 1024
MAX_ZIP_ENTRIES = 10_000
PARSER_VERSION = "pdf-docx-xlsx-v1"
SUPPORTED = {"pdf", "docx", "xlsx"}


def extension(filename):
    return str(filename).rsplit(".", 1)[-1].casefold() if "." in str(filename) else ""


def _zip_guard(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        members = archive.infolist()
        if len(members) > MAX_ZIP_ENTRIES or sum(item.file_size for item in members) > MAX_ZIP_UNCOMPRESSED:
            raise ValueError("expanded_zip_limit")
        if any(item.flag_bits & 1 for item in members):
            raise ValueError("encrypted_zip")


class Collector:
    def __init__(self):
        self.segments = []
        self.characters = 0
        self.partial = False

    def add(self, location, text):
        text = str(text or "").strip()
        if not text:
            return True
        remaining = MAX_TEXT - self.characters
        if remaining <= 0:
            self.partial = True
            return False
        if len(text) > remaining:
            text = text[:remaining]
            self.partial = True
        for start in range(0, len(text), 1620):
            piece = text[start:start + 1800]
            self.segments.append({"location": location, "text": piece})
        self.characters += len(text)
        return not self.partial


def extract_bytes(filename, data):
    kind = extension(filename)
    if kind not in SUPPORTED:
        return {"status": "unsupported", "reason": "format", "segments": [], "partial": False}
    if len(data) > MAX_INPUT:
        return {"status": "too_large", "reason": "input_bytes", "segments": [], "partial": False}
    result = Collector()
    try:
        if kind == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data), strict=False)
            if reader.is_encrypted:
                return {"status": "unsupported", "reason": "encrypted", "segments": [], "partial": False}
            for number, page in enumerate(reader.pages, 1):
                if number > 40:
                    result.partial = True
                    break
                if not result.add({"page": number}, page.extract_text() or ""):
                    break
        else:
            _zip_guard(data)
            if kind == "docx":
                from docx import Document
                document = Document(io.BytesIO(data))
                for number, paragraph in enumerate(document.paragraphs, 1):
                    if number > 2000:
                        result.partial = True
                        break
                    if not result.add({"paragraph": number}, paragraph.text):
                        break
                if not result.partial:
                    for table_number, table in enumerate(document.tables, 1):
                        if table_number > 20:
                            result.partial = True
                            break
                        for row_number, row in enumerate(table.rows, 1):
                            for cell_number, cell in enumerate(row.cells, 1):
                                if not result.add({"table": table_number, "row": row_number, "cell": cell_number}, cell.text):
                                    break
                            if result.partial:
                                break
                        if result.partial:
                            break
            else:
                # defusedxml is installed: openpyxl uses it for untrusted XML.
                import defusedxml  # noqa: F401
                from openpyxl import load_workbook
                workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
                try:
                    extra_sheets = len(workbook.worksheets) > 20
                    inspected = 0
                    for sheet in workbook.worksheets[:20]:
                        for row in sheet.iter_rows():
                            for cell in row:
                                inspected += 1
                                if inspected > 20_000:
                                    result.partial = True
                                    break
                                if cell.value is not None and not result.add({"sheet": sheet.title,
                                        "cell": cell.coordinate}, str(cell.value)[:1000]):
                                    break
                            if result.partial:
                                break
                        if result.partial:
                            break
                    result.partial = result.partial or extra_sheets
                finally:
                    workbook.close()
        if not result.segments:
            return {"status": "unsupported", "reason": "no_text_layer", "segments": [], "partial": result.partial}
        return {"status": "partial" if result.partial else "ready", "reason": "limit" if result.partial else "",
                "segments": result.segments, "partial": result.partial}
    except ValueError as exc:
        reason = str(exc)
        if reason == "expanded_zip_limit":
            return {"status": "too_large", "reason": "expanded_zip_bytes", "segments": [], "partial": False}
        return {"status": "unsupported" if reason == "encrypted_zip" else "failed",
                "reason": "encrypted" if reason == "encrypted_zip" else "parse_error",
                "segments": [], "partial": False}
    except (zipfile.BadZipFile, OSError, KeyError, TypeError):
        return {"status": "failed", "reason": "parse_error", "segments": [], "partial": False}


def main():
    if os.name == "posix":
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024, 512 * 1024 * 1024))
        resource.setrlimit(resource.RLIMIT_CPU, (15, 15))
    filename = sys.argv[1]
    data = sys.stdin.buffer.read(MAX_INPUT + 1)
    value = extract_bytes(filename, data)
    sys.stdout.write(json.dumps(value, ensure_ascii=False, separators=(",", ":")))


if __name__ == "__main__":
    main()
