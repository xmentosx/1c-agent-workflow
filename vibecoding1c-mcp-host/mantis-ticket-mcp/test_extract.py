from __future__ import annotations

import io
import json
import os
import subprocess
import sys
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

import mantis_extract


def pdf_with_text(value):
    stream = f"BT /F1 12 Tf 72 720 Td ({value}) Tj ET".encode("ascii")
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream"]
    content = b"%PDF-1.4\n"
    offsets = [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(content))
        content += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(content)
    content += f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode()
    content += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    content += f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return content


class ExtractionTests(unittest.TestCase):
    def test_pdf_docx_xlsx_coordinates_and_isolated_process(self):
        from docx import Document
        from openpyxl import Workbook
        doc = Document()
        doc.add_paragraph("Уникальная фраза документа")
        docx = io.BytesIO()
        doc.save(docx)
        workbook = Workbook()
        workbook.active.title = "План"
        workbook.active["B4"] = "Уникальная фраза таблицы"
        xlsx = io.BytesIO()
        workbook.save(xlsx)
        cases = [("sample.pdf", pdf_with_text("UniquePDFPhrase"), "UniquePDFPhrase", {"page": 1}),
                 ("sample.docx", docx.getvalue(), "Уникальная фраза документа", {"paragraph": 1}),
                 ("sample.xlsx", xlsx.getvalue(), "Уникальная фраза таблицы", {"sheet": "План", "cell": "B4"})]
        for name, data, phrase, location in cases:
            with self.subTest(name=name):
                child = subprocess.run([sys.executable, "-m", "mantis_extract", name], input=data,
                    stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=20, cwd=Path(__file__).parent,
                    env={**os.environ, "PYTHONIOENCODING": "utf-8", "OPENBLAS_NUM_THREADS": "12"})
                self.assertEqual(child.returncode, 0, child.stderr.decode(errors="replace"))
                result = json.loads(child.stdout.decode("utf-8"))
                self.assertEqual(result["status"], "ready")
                self.assertTrue(any(phrase in segment["text"] and segment["location"] == location
                                    for segment in result["segments"]))

    def test_macro_extension_and_expanded_zip_limit_are_rejected(self):
        self.assertEqual(mantis_extract.extract_bytes("active.xlsm", b"not-read")["status"], "unsupported")
        self.assertEqual(mantis_extract.extract_bytes("broken.docx", b"not-a-zip")["status"], "failed")
        data = io.BytesIO()
        with zipfile.ZipFile(data, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("large.xml", "a" * 200)
        with patch.object(mantis_extract, "MAX_ZIP_UNCOMPRESSED", 100):
            result = mantis_extract.extract_bytes("sample.xlsx", data.getvalue())
        self.assertEqual(result["status"], "too_large")
        self.assertEqual(result["reason"], "expanded_zip_bytes")

    def test_scan_without_text_layer_has_explicit_outcome(self):
        result = mantis_extract.extract_bytes("scan.pdf", pdf_with_text(""))
        self.assertEqual(result["status"], "unsupported")
        self.assertEqual(result["reason"], "no_text_layer")


if __name__ == "__main__":
    unittest.main()
